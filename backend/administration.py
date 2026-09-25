import csv, io
<<<<<<< HEAD
from datetime import datetime, timezone
from fastapi import APIRouter,Depends,HTTPException,Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from core import db,uid,now,authorize,project_scope,project_public,log_activity,FINANCE,MANAGERS,STATUSES,DEFAULT_CLIENT_PASSWORD
=======
from fastapi import APIRouter,Depends,HTTPException,Response
<<<<<<< HEAD
from core import db,uid,now,authorize,project_scope,project_public,log_activity,FINANCE,MANAGERS,STATUSES,DEFAULT_CLIENT_PASSWORD
=======
<<<<<<< HEAD
from core import db,uid,now,authorize,project_scope,project_public,FINANCE,MANAGERS,STATUSES,DEFAULT_CLIENT_PASSWORD
=======
from core import db,uid,now,authorize,project_scope,project_public,FINANCE,MANAGERS,STATUSES
>>>>>>> b246b9f0dcd59f93e220dafc66ccfd5b50d9cc1b
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
from auth import current_user,hash_password,public_user
from schemas import Record,ClientInput,UserInput,UserUpdate

router=APIRouter()
@router.get('/users',response_model=list[Record])
async def users(u=Depends(current_user)):
    await authorize(u,'user.manage')
    rows=await db.users.find({},{'_id':0,'password_hash':0}).to_list(1000)
    return rows
@router.get('/team',response_model=list[Record])
async def team(u=Depends(current_user)):
    await authorize(u,'team.read')
    return await db.users.find({'active':True,'role':'Developer'},{'_id':0,'id':1,'name':1,'role':1}).to_list(500)
@router.post('/users',response_model=Record)
async def add_user(data:UserInput,u=Depends(current_user)):
    await authorize(u,'user.manage')
    if await db.users.find_one({'username':data.username.lower()},{'_id':0}): raise HTTPException(409,'Username sudah digunakan.')
    if data.role=='Client' and not await db.clients.find_one({'id':data.client_id},{'_id':0}): raise HTTPException(400,'Akun Client harus terhubung ke client yang valid.')
    row={k:v for k,v in data.model_dump().items() if k!='password'}
    row.update(id=uid(),username=data.username.lower(),password_hash=hash_password(data.password),active=True,created_at=now())
    try: await db.users.insert_one(row.copy())
    except Exception: raise HTTPException(409,'Username sudah digunakan.')
<<<<<<< HEAD
    await log_activity(u,'buat','user',row['id'],row['name'],'',{'role':row['role']})
=======
<<<<<<< HEAD
    await log_activity(u,'buat','user',row['id'],row['name'],'',{'role':row['role']})
=======
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
    return public_user(row)
@router.patch('/users/{user_id}',response_model=Record)
async def edit_user(user_id:str,data:UserUpdate,u=Depends(current_user)):
    await authorize(u,'user.manage')
    row=await db.users.find_one({'id':user_id},{'_id':0})
    if not row: raise HTTPException(404,'User tidak ditemukan.')
    updates=data.model_dump(exclude_none=True)
    new_password=updates.pop('new_password',None)
    if new_password:
        updates['password_hash']=hash_password(new_password)
        await db.sessions.delete_many({'user_id':user_id})
    if user_id==u['id'] and (updates.get('active') is False or updates.get('role',u['role'])!='Admin'): raise HTTPException(400,'Anda tidak dapat menonaktifkan atau menurunkan role akun sendiri.')
    merged={**row,**updates}
    if merged['role']=='Client' and not await db.clients.find_one({'id':merged.get('client_id','')},{'_id':0}): raise HTTPException(400,'Pilih client untuk akun ini.')
    await db.users.update_one({'id':user_id},{'$set':updates})
<<<<<<< HEAD
    await log_activity(u,'ubah','user',user_id,row['name'],'',{k:v for k,v in updates.items() if k!='password_hash'}|({'reset_password':True} if new_password else {}))
=======
<<<<<<< HEAD
    await log_activity(u,'ubah','user',user_id,row['name'],'',{k:v for k,v in updates.items() if k!='password_hash'}|({'reset_password':True} if new_password else {}))
=======
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
    if updates.get('active') is False: await db.sessions.delete_many({'user_id':user_id})
    return public_user(merged)

@router.get('/clients',response_model=list[Record])
async def clients(u=Depends(current_user)):
    await authorize(u,'client.read')
    rows=await db.clients.find({},{'_id':0}).sort('created_at',-1).to_list(2000)
    for r in rows: r['project_count']=await db.projects.count_documents({'client_id':r['id']})
    return rows
@router.post('/clients',response_model=Record)
async def add_client(data:ClientInput,u=Depends(current_user)):
    await authorize(u,'client.write')
    row={**data.model_dump(),'id':uid(),'created_at':now()}
    await db.clients.insert_one(row.copy())
<<<<<<< HEAD
    await log_activity(u,'buat','client',row['id'],row['name'])
=======
<<<<<<< HEAD
    await log_activity(u,'buat','client',row['id'],row['name'])
=======
<<<<<<< HEAD
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
    username=data.email.lower()
    account={'username':username,'created':False}
    if not await db.users.find_one({'username':username},{'_id':0,'id':1}):
        await db.users.insert_one({'id':uid(),'username':username,'name':data.contact,'role':'Client','email':data.email,'client_id':row['id'],'password_hash':hash_password(DEFAULT_CLIENT_PASSWORD),'active':True,'must_change_password':True,'created_at':now()})
        account.update(created=True,password=DEFAULT_CLIENT_PASSWORD)
    return {**row,'account':account}
<<<<<<< HEAD
=======
<<<<<<< HEAD
=======
=======
    return row
>>>>>>> b246b9f0dcd59f93e220dafc66ccfd5b50d9cc1b
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
@router.patch('/clients/{cid}',response_model=Record)
async def edit_client(cid:str,data:ClientInput,u=Depends(current_user)):
    await authorize(u,'client.write')
    r=await db.clients.update_one({'id':cid},{'$set':data.model_dump()})
    if not r.matched_count: raise HTTPException(404,'Client tidak ditemukan.')
    await db.projects.update_many({'client_id':cid},{'$set':{'client_name':data.name}})
    return await db.clients.find_one({'id':cid},{'_id':0})
@router.delete('/clients/{cid}')
async def delete_client(cid:str,u=Depends(current_user)):
    await authorize(u,'client.write')
    if await db.projects.count_documents({'client_id':cid}) or await db.users.count_documents({'client_id':cid}): raise HTTPException(400,'Client masih terhubung dengan project atau akun.')
    r=await db.clients.delete_one({'id':cid})
    if not r.deleted_count: raise HTTPException(404,'Client tidak ditemukan.')
<<<<<<< HEAD
    await log_activity(u,'hapus','client',cid,'')
=======
<<<<<<< HEAD
    await log_activity(u,'hapus','client',cid,'')
=======
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
    return {'message':'Client dihapus.'}

@router.get('/dashboard')
async def dashboard(u=Depends(current_user)):
    await authorize(u,'dashboard.read')
    rows=await db.projects.find(project_scope(u),{'_id':0}).sort('updated_at',-1).to_list(2000)
    ids=[p['id'] for p in rows]
    status_counts={s:sum(p['status']==s for p in rows) for s in STATUSES}
    tq={'project_id':{'$in':ids}}
    if u['role']=='Developer': tq['assigned_to']=u['id']
    result={'total':len(rows),'active':sum(p['status']!='Selesai' for p in rows),'completed':status_counts['Selesai'],'development':status_counts['Development'],'dev_server':status_counts['Uploaded to Dev Server'],'production':status_counts['Uploaded to Production'],'status_counts':status_counts,'projects':[project_public(p,u) for p in rows[:5]],'active_revisions':0,'active_maintenance':0,'open_tickets':0,'closed_tickets':0}
    if u['role']!='Accounting':
        result['open_tickets']=await db.tickets.count_documents({**tq,'status':{'$nin':['Selesai','Ditutup','Ditolak']}})
        result['closed_tickets']=await db.tickets.count_documents({**tq,'status':{'$in':['Selesai','Ditutup']}})
    if u['role'] in MANAGERS+['Developer']:
        result['active_revisions']=await db.revisions.count_documents({'project_id':{'$in':ids},'status':{'$ne':'Selesai'}})
        result['active_maintenance']=await db.maintenances.count_documents({'project_id':{'$in':ids},'status':{'$ne':'Selesai'}})
    if u['role'] in FINANCE:
<<<<<<< HEAD
        result['finance']={k:sum(p.get(k,0) for p in rows) for k in ['value','development_cost','server_cost','other_cost']}
        result['finance']['profit']=result['finance']['value']-result['finance']['development_cost']-result['finance']['server_cost']-result['finance']['other_cost']
=======
<<<<<<< HEAD
        result['finance']={k:sum(p.get(k,0) for p in rows) for k in ['value','development_cost','server_cost','other_cost']}
        result['finance']['profit']=result['finance']['value']-result['finance']['development_cost']-result['finance']['server_cost']-result['finance']['other_cost']
=======
<<<<<<< HEAD
        result['finance']={k:sum(p.get(k,0) for p in rows) for k in ['value','development_cost','server_cost','other_cost']}
        result['finance']['profit']=result['finance']['value']-result['finance']['development_cost']-result['finance']['server_cost']-result['finance']['other_cost']
=======
        result['finance']={k:sum(p.get(k,0) for p in rows) for k in ['value','development_cost','server_cost']}
        result['finance']['profit']=result['finance']['value']-result['finance']['development_cost']-result['finance']['server_cost']
>>>>>>> b246b9f0dcd59f93e220dafc66ccfd5b50d9cc1b
>>>>>>> abb7b4276f614088423817f4dc75effc576b2e5b
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
    elif u['role']=='Admin Project': result['total_value']=sum(p.get('value',0) for p in rows)
    histories=await db.project_status_logs.find({'project_id':{'$in':ids}},{'_id':0}).sort('created_at',-1).to_list(30)
    names={p['id']:p['name'] for p in rows}
    if u['role']=='Client': histories=[{k:r[k] for k in ['id','project_id','message','created_at'] if k in r} for r in histories if r.get('to_status')]
    result['activity']=[{**r,'project_name':names.get(r['project_id'],'')} for r in histories[:5]]
    result['deadlines']=[project_public(p,u) for p in sorted(rows,key=lambda p:p['due_date']) if p['status']!='Selesai'][:4]
    return result

@router.get('/reports/projects.csv')
async def report(u=Depends(current_user)):
    await authorize(u,'project.read')
    rows=await db.projects.find(project_scope(u),{'_id':0}).to_list(2000)
    fields=['Kode','Project','Client','Status','Progress','Deadline']
    if u['role'] in FINANCE: fields+=['Nilai Project','Biaya Development','Biaya Server','Profit']
    stream=io.StringIO();writer=csv.writer(stream);writer.writerow(fields)
    def safe(v):
        s=str(v)
        return "'"+s if s.startswith(('=','+','-','@')) else s
    for p in rows:
        values=[p['code'],p['name'],p['client_name'],p['status'],p['progress'],p['due_date']]
        if u['role'] in FINANCE: values += [p['value'],p.get('development_cost',0),p.get('server_cost',0),p['value']-p.get('development_cost',0)-p.get('server_cost',0)]
        writer.writerow([safe(v) for v in values])
<<<<<<< HEAD
    return Response('\ufeff'+stream.getvalue(),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename=laporan-project-maiharta.csv'})

@router.get('/reports/projects.xlsx')
async def report_xlsx(u=Depends(current_user)):
    await authorize(u,'project.read')
    rows=await db.projects.find(project_scope(u),{'_id':0}).sort('code',1).to_list(2000)
    is_finance=u['role'] in FINANCE
    headers=['Kode','Project','Client','Status','Progress','Deadline']
    if is_finance: headers+=['Nilai Project','Biaya Development','Biaya Server','Profit']
    ncols=len(headers); last=get_column_letter(ncols)
    NAVY='1E3A5F'; BLUE='2F5FE0'; LIGHT='EEF3FF'; BAND='F7F9FE'; LINE='D9E1F2'
    thin=Side(style='thin', color=LINE)
    border=Border(left=thin,right=thin,top=thin,bottom=thin)
    money_fmt='"Rp"#,##0'
    wb=Workbook(); ws=wb.active; ws.title='Laporan Project'
    # Title band
    ws.merge_cells(f'A1:{last}1')
    c=ws['A1']; c.value='Laporan Project — CRM Maiharta'
    c.font=Font(name='Calibri',size=16,bold=True,color='FFFFFF')
    c.fill=PatternFill('solid',fgColor=NAVY)
    c.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[1].height=34
    ws.merge_cells(f'A2:{last}2')
    s=ws['A2']
    s.value=f"Dibuat {datetime.now(timezone.utc).strftime('%d %b %Y')}  ·  {len(rows)} project"
    s.font=Font(name='Calibri',size=10,italic=True,color='FFFFFF')
    s.fill=PatternFill('solid',fgColor=BLUE)
    s.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[2].height=20
    # Header row
    hr=4
    for i,h in enumerate(headers, start=1):
        cell=ws.cell(row=hr,column=i,value=h)
        cell.font=Font(bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor=BLUE)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        cell.border=border
    ws.row_dimensions[hr].height=26
    # Data rows
    r=hr+1
    for idx,p in enumerate(rows):
        prof=p['value']-p.get('development_cost',0)-p.get('server_cost',0)
        vals=[p['code'],p['name'],p['client_name'],p['status'],(p.get('progress',0) or 0)/100,p['due_date']]
        if is_finance: vals+=[p['value'],p.get('development_cost',0),p.get('server_cost',0),prof]
        for ci,v in enumerate(vals, start=1):
            cell=ws.cell(row=r,column=ci,value=v)
            cell.border=border
            cell.alignment=Alignment(vertical='center',horizontal='center' if ci in (1,4,5,6) else 'left')
            if idx%2: cell.fill=PatternFill('solid',fgColor=BAND)
        pc=ws.cell(row=r,column=5); pc.number_format='0%'; pc.alignment=Alignment(horizontal='center',vertical='center')
        if is_finance:
            for ci in range(7,ncols+1):
                mc=ws.cell(row=r,column=ci); mc.number_format=money_fmt
                mc.alignment=Alignment(horizontal='right',vertical='center')
        r+=1
    # Totals
    if is_finance and rows:
        tr=r
        lc=ws.cell(row=tr,column=6,value='Total'); lc.font=Font(bold=True)
        lc.alignment=Alignment(horizontal='right',vertical='center')
        for ci in range(1,7):
            ws.cell(row=tr,column=ci).fill=PatternFill('solid',fgColor=LIGHT)
            ws.cell(row=tr,column=ci).border=border
        for ci in range(7,ncols+1):
            col=get_column_letter(ci)
            tc=ws.cell(row=tr,column=ci,value=f'=SUM({col}{hr+1}:{col}{r-1})')
            tc.number_format=money_fmt; tc.font=Font(bold=True,color=NAVY)
            tc.fill=PatternFill('solid',fgColor=LIGHT); tc.border=border
            tc.alignment=Alignment(horizontal='right',vertical='center')
        ws.row_dimensions[tr].height=22
    # Widths + view
    widths=[12,32,24,22,11,14]
    if is_finance: widths+=[16,18,15,16]
    for ci,w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width=w
    ws.freeze_panes=f'A{hr+1}'
    ws.sheet_view.showGridLines=False
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition':'attachment; filename=laporan-project-maiharta.xlsx'})
=======
    return Response('\ufeff'+stream.getvalue(),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename=laporan-project-maiharta.csv'})
>>>>>>> 0641a06cd4d77fcddb9db1930ccf186521511ad2
